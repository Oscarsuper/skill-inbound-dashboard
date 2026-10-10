# -*- coding: utf-8 -*-
"""Arma la plantilla que el Centro de Mando sabe leer.

Los nombres de las hojas y de las columnas no son decorativos: el portal
los busca con esos mismos patrones. Lo que sí es libre es QUÉ indicadores
se midan y cuántos: cada bloque de tres columnas de la fila 2 se lee solo.
"""
import re, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PORTAL = '/home/user/skill-inbound-dashboard/equipo/Portal_Mando.html'
SALIDA = '/home/user/skill-inbound-dashboard/plantilla/Plantilla_Resultados.xlsx'

# La nómina se toma del portal para que no haya dos listas que mantener
h = io.open(PORTAL, encoding='utf-8').read()
AGENTES = re.findall(r"ced:'(\d+)',\s*nombre:'([^']+)',\s*canal:'(\w+)'", h)
assert len(AGENTES) == 21, len(AGENTES)

# ── Paleta, en claro para que imprima bien ───────────────────────────────
TINTA   = '1E1B4B'
CAB     = PatternFill('solid', fgColor='2B0B6B')
CAB2    = PatternFill('solid', fgColor='4A1D96')
SUAVE   = PatternFill('solid', fgColor='EDE9FE')
BANDA   = PatternFill('solid', fgColor='F7F5FF')
EDITA   = PatternFill('solid', fgColor='FFFDF2')   # lo que él diligencia
CALC    = PatternFill('solid', fgColor='F1F5F9')   # lo que sale solo
BLANCO  = Font(color='FFFFFF', bold=True, size=9.5)
TIT     = Font(color=TINTA, bold=True, size=11)
NORMAL  = Font(color='334155', size=10)
fino    = Side(style='thin', color='D7D0F5')
BORDE   = Border(left=fino, right=fino, top=fino, bottom=fino)

# ── Indicadores por defecto ──────────────────────────────────────────────
# tipo: 'pct' se escribe en porcentaje · 'tiempo' se escribe mm:ss
# dir : 'mayor' cumple por encima de la meta · 'menor' por debajo
INDICADORES = [
    ('ADHERENCIA',     'pct',    0.95,  'mayor'),
    ('CALIDAD',        'pct',    0.90,  'mayor'),
    ('ENCUESTAS',      'pct',    0.80,  'mayor'),
    ('TIPIFICACION',   'pct',    0.98,  'mayor'),
    ('TMO',            'tiempo', '3:00','menor'),
    ('TMR',            'tiempo', '0:30','menor'),
    ('AUSENTISMO',     'pct',    0.03,  'menor'),
]
PRODUCTIVIDAD = ['contestados', 'promedio_diario']

CRITERIOS = ['Saludo institucional', 'Verificación de identidad', 'Escucha activa',
             'Manejo del aplicativo', 'Información entregada', 'Tipificación de la gestión',
             'Cierre y despedida', 'Compromiso de mejora']

wb = Workbook(); wb.remove(wb.active)

# ═══════════════════════════════════════════════════════════════════════
# 1. INSTRUCCIONES
# ═══════════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Instrucciones')
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 104

def linea(texto, estilo='p', salto=1):
    global fila
    c = ws.cell(fila, 2, texto)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if estilo == 'h1':
        c.font = Font(color='FFFFFF', bold=True, size=15)
        c.fill = CAB; c.alignment = Alignment(vertical='center', indent=1)
        ws.row_dimensions[fila].height = 34
    elif estilo == 'h2':
        c.font = Font(color='2B0B6B', bold=True, size=11.5)
        c.fill = SUAVE; c.alignment = Alignment(vertical='center', indent=1)
        ws.row_dimensions[fila].height = 24
    elif estilo == 'ojo':
        c.font = Font(color='9A3412', bold=True, size=10)
        c.fill = PatternFill('solid', fgColor='FEF3C7')
        c.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
        ws.row_dimensions[fila].height = 15 * (1 + texto.count('\n'))
    else:
        c.font = NORMAL
        ws.row_dimensions[fila].height = max(15, 13.5 * (1 + texto.count('\n')))
    fila += salto

fila = 2
linea('PLANTILLA DE RESULTADOS — Chat y RRSS', 'h1', 2)
linea('Este es el único archivo que tienes que diligenciar. Cuando esté listo, lo cargas en el '
      'Centro de Mando con el botón «Cargar Excel» y de ahí sale el portal que les repartes a los '
      'agentes.\n\nNo hay que cargar nada más, salvo la malla de turnos, que llega por aparte '
      'desde WFM y se carga tal como viene.', 'p', 2)

linea('LAS HOJAS', 'h2')
linea('· Indicadores_Chat y Indicadores_RRSS — los resultados del mes de cada agente.\n'
      '· Pendientes_Tipificar — la cola de interacciones sin tipificar.\n'
      '· Lado a Lado — los acompañamientos, con la observación de cada criterio.', 'p', 2)
linea('No les cambies el nombre a las hojas. El portal las busca por el nombre; si una se llama '
      'distinto, simplemente no la lee.', 'ojo', 2)

linea('CÓMO SE LEEN LOS INDICADORES', 'h2')
linea('En las hojas de indicadores, la fila 1 tiene el nombre del indicador y la fila 2 sus tres '
      'columnas: RESULTADO, META y CUMPLIMIENTO. A partir de la fila 3 va un agente por fila.\n\n'
      'Puedes agregar indicadores nuevos o quitar los que no uses: basta con respetar esa '
      'estructura de tres columnas. El portal los descubre solos, no hay nada que configurarle.', 'p', 2)

linea('ESTO ES LO MÁS IMPORTANTE DE TODO EL ARCHIVO', 'h2')
linea('El portal decide si alguien cumple comparando el RESULTADO con la META. Para saber hacia '
      'qué lado, mira el NOMBRE del indicador:\n\n'
      '· Si el nombre contiene TMO, TMR, AHT, TIEMPO, AUSENTISMO, ABANDONO, DESCONEXIÓN, ERROR, '
      'MAL RADICADO, MORA o ESPERA, entonces cumplir es quedar POR DEBAJO de la meta.\n'
      '· En todos los demás, cumplir es quedar POR ENCIMA.\n\n'
      'Por eso, si creas un indicador donde menos es mejor, ponle un nombre que lleve una de esas '
      'palabras. Un indicador llamado «REPROCESOS» se leería al revés; «ERRORES DE REPROCESO» se '
      'lee bien.', 'ojo', 2)

linea('CÓMO SE ESCRIBE CADA DATO', 'h2')
linea('· Porcentajes: escribe 95% o 0,95. La celda ya viene con formato de porcentaje.\n'
      '· Tiempos (TMO, TMR): escribe minutos y segundos, como 3:05. Esas celdas vienen en formato '
      'de texto a propósito, para que Excel no te las convierta en hora del día.\n'
      '· Cantidades: números enteros, sin puntos de miles.\n'
      '· Si un agente no tiene medición en un indicador, deja la celda vacía. El portal lo marca '
      'como no cumple y se nota a simple vista.\n'
      '· La columna CUMPLIMIENTO se calcula sola. Es para que tú lo veas aquí; el portal no la usa '
      'para decidir, él vuelve a comparar resultado contra meta.', 'p', 2)

linea('PENDIENTES POR TIPIFICAR', 'h2')
linea('El nombre del agente tiene que quedar escrito IGUAL que en la hoja de indicadores: por ahí '
      'se cruzan las dos hojas. Si en una dice «PEREZ GOMEZ ANA» y en la otra «Ana Pérez Gómez», '
      'el portal no las une.\n\n'
      'En la columna del enlace pega la URL completa de la interacción. El portal le pone el botón '
      '«Ver» al agente y la abre en una pestaña nueva.', 'p', 2)

linea('LADO A LADO', 'h2')
linea('Una fila por acompañamiento. Si a la misma persona le haces tres, van tres filas con su '
      'nombre repetido: el portal se las muestra ordenadas, de la más reciente a la más antigua.\n\n'
      'No se califica con nota: lo que vale es lo que escribas en cada criterio. Un criterio en '
      'blanco se muestra como «sin observación».\n\n'
      'Las columnas de criterio son libres. Agrega, quita o renombra las que quieras; el portal '
      'muestra las que encuentre. En la columna del ID puedes pegar el enlace completo de la '
      'interacción y el agente tendrá el botón para abrirla.', 'p', 2)

linea('UN DETALLE SOBRE PRIVACIDAD', 'h2')
linea('El HTML que generas lleva los datos adentro. Cada agente ve lo suyo, pero alguien que sepa '
      'abrir las herramientas del navegador podría leer lo de sus compañeros. Para indicadores y '
      'pendientes del día a día no es grave; si algún dato fuera reservado, no lo incluyas.', 'p', 2)
ws.freeze_panes = 'A2'
print('Instrucciones: %d filas' % fila)

# ═══════════════════════════════════════════════════════════════════════
# 2. INDICADORES (una hoja por canal)
# ═══════════════════════════════════════════════════════════════════════
def hoja_indicadores(nombre, gente):
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False

    # Fila 1: el nombre del indicador, combinado sobre sus tres columnas.
    # Fila 2: los subtítulos que el portal busca.
    ws.cell(1, 1, 'AGENTES').fill = CAB
    ws.cell(1, 1).font = BLANCO
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws.cell(1, 2).fill = CAB
    ws.cell(2, 1, 'cedula'); ws.cell(2, 2, 'nombre_completo')

    col = 3
    sitio = {}                      # dónde quedó cada indicador
    for i, (nm, tipo, meta, direc) in enumerate(INDICADORES):
        c1 = ws.cell(1, col, nm)
        c1.fill = CAB if i % 2 == 0 else CAB2
        c1.font = BLANCO
        c1.alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 2)
        for k in (1, 2):
            ws.cell(1, col + k).fill = CAB if i % 2 == 0 else CAB2
        ws.cell(2, col, 'RESULTADO'); ws.cell(2, col + 1, 'META')
        ws.cell(2, col + 2, 'CUMPLIMIENTO')
        sitio[nm] = (col, tipo, meta, direc)
        col += 3

    # Bloque sin meta: el portal lo lista tal cual, sin juzgarlo
    cp = col
    c1 = ws.cell(1, cp, 'PRODUCTIVIDAD')
    c1.fill = PatternFill('solid', fgColor='0F766E'); c1.font = BLANCO
    c1.alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=1, start_column=cp, end_row=1, end_column=cp + len(PRODUCTIVIDAD) - 1)
    for k, p in enumerate(PRODUCTIVIDAD):
        ws.cell(1, cp + k).fill = PatternFill('solid', fgColor='0F766E')
        ws.cell(2, cp + k, p)
    ultima = cp + len(PRODUCTIVIDAD) - 1

    for c in range(1, ultima + 1):
        s = ws.cell(2, c)
        s.fill = SUAVE
        s.font = Font(color=TINTA, bold=True, size=8.5)
        s.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        s.border = BORDE
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 30

    # Los agentes ya vienen escritos: solo hay que llenarles los números
    for i, (ced, nom, _) in enumerate(gente):
        r = 3 + i
        a = ws.cell(r, 1, ced); a.font = Font(size=9.5, color='475569')
        a.alignment = Alignment(horizontal='left')
        a.number_format = '@'
        b = ws.cell(r, 2, nom); b.font = Font(size=10, bold=True, color=TINTA)
        for c in range(1, ultima + 1):
            cel = ws.cell(r, c)
            cel.border = BORDE
            if c > 2 and i % 2 == 1:
                cel.fill = BANDA

        for nm, (col0, tipo, meta, direc) in sitio.items():
            res, met, cum = (get_column_letter(col0) + str(r),
                             get_column_letter(col0 + 1) + str(r),
                             get_column_letter(col0 + 2) + str(r))
            cr, cm, cc = ws[res], ws[met], ws[cum]
            cr.fill = EDITA
            cr.alignment = cm.alignment = cc.alignment = Alignment(horizontal='center')
            if tipo == 'pct':
                cr.number_format = cm.number_format = '0.0%'
                cm.value = meta
                cc.value = ('=IF({r}="","",IF({r}{op}{m},"CUMPLE","NO CUMPLE"))'
                            .format(r=res, m=met, op='>=' if direc == 'mayor' else '<='))
            else:
                # texto a propósito: si fuera formato de hora, Excel lo
                # convertiría en hora del día y el portal leería 0:00
                cr.number_format = cm.number_format = '@'
                cm.value = meta
                cc.value = ('=IF({r}="","",IF(IFERROR(TIMEVALUE({r}),99)'
                            '<=IFERROR(TIMEVALUE({m}),0),"CUMPLE","NO CUMPLE"))'
                            .format(r=res, m=met))
            cm.font = Font(size=9.5, color='64748B')
            cc.fill = CALC
            cc.font = Font(size=9, bold=True, color='475569')
        for k in range(len(PRODUCTIVIDAD)):
            cel = ws.cell(r, cp + k)
            cel.fill = EDITA; cel.number_format = '#,##0'
            cel.alignment = Alignment(horizontal='center')

    ws.column_dimensions['A'].width = 13
    ws.column_dimensions['B'].width = 42
    for c in range(3, ultima + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.freeze_panes = 'C3'
    return ws

chat = [a for a in AGENTES if a[2] == 'Chat']
rrss = [a for a in AGENTES if a[2] == 'RRSS']
hoja_indicadores('Indicadores_Chat', chat)
hoja_indicadores('Indicadores_RRSS', rrss)
print('Indicadores: %d chat, %d rrss' % (len(chat), len(rrss)))

# ═══════════════════════════════════════════════════════════════════════
# 3. PENDIENTES POR TIPIFICAR
# ═══════════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Pendientes_Tipificar')
ws.sheet_view.showGridLines = False
COLS_P = [('nombre_agente', 40), ('id_interaccion', 34), ('link_gestion', 46), ('cola', 14),
          ('macroproceso', 24), ('proceso', 24), ('subproceso_uno', 24),
          ('subproceso_dos', 24), ('observaciones', 34)]
for i, (nm, w) in enumerate(COLS_P, start=1):
    c = ws.cell(1, i, nm)
    c.fill = CAB; c.font = BLANCO
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = BORDE
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 28
ws.freeze_panes = 'A2'
ws.auto_filter.ref = 'A1:%s1' % get_column_letter(len(COLS_P))
# Lista desplegable con los nombres, para que no haya dos formas de escribirlos
nombres = ','.join(n for _, n, _ in AGENTES)
if len(nombres) < 250:
    dv = DataValidation(type='list', formula1='"%s"' % nombres, allow_blank=True)
    ws.add_data_validation(dv); dv.add('A2:A2000')
for r in range(2, 400):
    for c in range(1, len(COLS_P) + 1):
        cel = ws.cell(r, c); cel.border = BORDE
        if r % 2 == 1: cel.fill = BANDA

# ═══════════════════════════════════════════════════════════════════════
# 4. LADO A LADO
# ═══════════════════════════════════════════════════════════════════════
ws = wb.create_sheet('Lado a Lado')
ws.sheet_view.showGridLines = False
cab = ['Nombre del agente', 'Fecha', 'ID o enlace de la interaccion'] + CRITERIOS
for i, nm in enumerate(cab, start=1):
    c = ws.cell(1, i, nm)
    c.fill = CAB if i <= 3 else CAB2
    c.font = BLANCO
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = BORDE
    ws.column_dimensions[get_column_letter(i)].width = 40 if i == 1 else (
        13 if i == 2 else (40 if i == 3 else 34))
ws.row_dimensions[1].height = 38
ws.freeze_panes = 'D2'
ws.auto_filter.ref = 'A1:%s1' % get_column_letter(len(cab))
if len(nombres) < 250:
    dv2 = DataValidation(type='list', formula1='"%s"' % nombres, allow_blank=True)
    ws.add_data_validation(dv2); dv2.add('A2:A2000')
for r in range(2, 300):
    for c in range(1, len(cab) + 1):
        cel = ws.cell(r, c); cel.border = BORDE
        cel.alignment = Alignment(wrap_text=True, vertical='top')
        if r % 2 == 1: cel.fill = BANDA
        if c == 2: cel.number_format = 'DD/MM/YYYY'

wb.save(SALIDA)
print('✔ %s' % SALIDA)
