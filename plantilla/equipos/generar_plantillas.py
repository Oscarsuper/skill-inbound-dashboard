# -*- coding: utf-8 -*-
"""Arma las plantillas de los tres equipos.

El grupo de Inbound comparte un solo archivo —cuatro supervisores dentro—
porque el portal que se reparte es uno para todos. Olga Lucía y Monitoreo
llevan el suyo aparte, que es como manejan su equipo.
"""
import re, io, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE=os.path.dirname(os.path.abspath(__file__))
PORTAL='/home/user/skill-inbound-dashboard/equipo/Portal_Equipos.html'
SALIDA='/home/user/skill-inbound-dashboard/plantilla/equipos'
os.makedirs(SALIDA, exist_ok=True)

h=io.open(PORTAL,encoding='utf-8').read()
AG=re.findall(r"ced:'(\d+)', nombre:'([^']+)', skill:'([^']*)', jefe:'([^']+)'", h)
assert len(AG)==109, len(AG)

GRUPO_INB={'FONSECA GONZALEZ DAVID ALEXANDER','GOMEZ PEREZ JOSE DANIEL',
           'ORTIZ CODINA JESUS ANTONIO','TROCHA MORALES EUDYMAR'}
EQUIPOS=[
 ('INBOUND',   'Plantilla_INBOUND.xlsx',
  lambda j: j in GRUPO_INB,
  'Archivo compartido por David Fonseca, Daniel Gómez, Jesús Ortiz y Eudymar Trocha. '
  'El portal que se genera es uno solo para todo el grupo.'),
 ('MONITOREO TRANSACCIONAL', 'Plantilla_MONITOREO_TRANSACCIONAL.xlsx',
  lambda j: j=='DOTOR SANTIAGO ALEX ALEJANDRO', 'Equipo de Alex Dotor.'),
 ('OLGA LUCIA','Plantilla_OLGA_LUCIA.xlsx',
  lambda j: j=='ORTIZ QUITIAN OLGA LUCIA', 'Equipo de Olga Lucía Ortiz.'),
]

# nombre, tipo, meta, dirección
IND=[('ADHERENCIA','pct',0.95,'mayor'), ('CALIDAD','pct',0.90,'mayor'),
     ('ENCUESTAS','pct',0.80,'mayor'),  ('TIPIFICACION','pct',0.98,'mayor'),
     ('NIVEL DE SERVICIO','pct',0.90,'mayor'),
     ('TMO','seg',410,'menor'),         ('ABANDONO','pct',0.05,'menor'),
     ('AUSENTISMO','pct',0.03,'menor'), ('MAL RADICADOS','pct',0.02,'menor'),
     ('MORA','pct',0.05,'menor')]
PROD=['contestados','promedio_diario']
CRIT=['Saludo institucional','Verificación de identidad','Escucha activa',
      'Manejo del aplicativo','Información entregada','Tipificación de la gestión',
      'Cierre y despedida','Compromiso de mejora']

TINTA='14323F'
CAB  =PatternFill('solid', fgColor='14323F')
CAB2 =PatternFill('solid', fgColor='0E7C66')
PROD_F=PatternFill('solid', fgColor='1B4F7E')
SUAVE=PatternFill('solid', fgColor='E8F1EE')
BANDA=PatternFill('solid', fgColor='F7FAFC')
EDITA=PatternFill('solid', fgColor='FFFDF4')
CALC =PatternFill('solid', fgColor='F1F5F9')
BLANCO=Font(color='FFFFFF', bold=True, size=9.5)
fino=Side(style='thin', color='DCE5EC')
BORDE=Border(left=fino,right=fino,top=fino,bottom=fino)

def construir(nombre_eq, archivo, pertenece, descripcion):
    gente=[a for a in AG if pertenece(a[3])]
    jefes=sorted({a[3] for a in gente})
    wb=Workbook(); wb.remove(wb.active)

    # ── Instrucciones ──────────────────────────────────────────────────
    ws=wb.create_sheet('Instrucciones'); ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=3; ws.column_dimensions['B'].width=104
    f=[2]
    def linea(t, e='p', salto=1):
        c=ws.cell(f[0],2,t); c.alignment=Alignment(wrap_text=True, vertical='top')
        if e=='h1':
            c.font=Font(color='FFFFFF',bold=True,size=15); c.fill=CAB
            c.alignment=Alignment(vertical='center', indent=1); ws.row_dimensions[f[0]].height=34
        elif e=='h2':
            c.font=Font(color=TINTA,bold=True,size=11.5); c.fill=SUAVE
            c.alignment=Alignment(vertical='center', indent=1); ws.row_dimensions[f[0]].height=24
        elif e=='ojo':
            c.font=Font(color='7C4A08',bold=True,size=10)
            c.fill=PatternFill('solid',fgColor='FEF6E7')
            c.alignment=Alignment(wrap_text=True, vertical='center', indent=1)
            ws.row_dimensions[f[0]].height=15*(1+t.count('\n'))
        else:
            c.font=Font(color='3D566E', size=10)
            ws.row_dimensions[f[0]].height=max(15, 13.5*(1+t.count('\n')))
        f[0]+=salto

    linea('PLANTILLA DE RESULTADOS — '+nombre_eq, 'h1', 2)
    linea(descripcion+'\n\nDiligencia este archivo y cárgalo en el tablero con el botón '
          '«Cargar Excel». De ahí sale el portal que les repartes a los agentes.', 'p', 2)
    linea('LAS HOJAS', 'h2')
    linea('· Indicadores — los resultados del mes, un agente por fila.\n'
          '· Pendientes_Tipificar — la cola de interacciones sin tipificar.\n'
          '· Lado a Lado — los acompañamientos, con la observación de cada criterio.', 'p', 2)
    linea('No les cambies el nombre a las hojas. El tablero las busca por el nombre; si una se '
          'llama distinto, no la lee y no avisa.', 'ojo', 2)
    linea('CÓMO SE LEEN LOS INDICADORES', 'h2')
    linea('La fila 1 trae el nombre del indicador y la fila 2 sus tres columnas: RESULTADO, '
          'META y CUMPLIMIENTO. Desde la fila 3 va un agente por fila.\n\n'
          'Puedes agregar indicadores o quitar los que no uses: basta con respetar ese bloque '
          'de tres columnas. El tablero los descubre solos.\n\n'
          'Las columnas de cédula, nombre, skill y jefe_inmediato no se tocan: con ellas el '
          'tablero sabe a quién pertenece cada fila y de qué grupo es.', 'p', 2)
    linea('ESTO ES LO MÁS IMPORTANTE DE TODO EL ARCHIVO', 'h2')
    linea('El tablero decide si alguien cumple comparando el RESULTADO con la META. Para saber '
          'hacia qué lado, mira el NOMBRE del indicador:\n\n'
          '· Si contiene TMO, TMR, AHT, TIEMPO, AUSENTISMO, ABANDONO, DESCONEXIÓN, ERROR, '
          'MAL RADICADO, MORA o ESPERA, cumplir es quedar POR DEBAJO de la meta.\n'
          '· En todos los demás, cumplir es quedar POR ENCIMA.\n\n'
          'Si creas un indicador donde menos es mejor, ponle un nombre con alguna de esas '
          'palabras. «REPROCESOS» se leería al revés; «ERRORES DE REPROCESO» se lee bien.', 'ojo', 2)
    linea('CÓMO SE ESCRIBE CADA DATO', 'h2')
    linea('· Porcentajes: escribe 95% o 0,95. La celda ya viene con formato de porcentaje.\n'
          '· TMO: en SEGUNDOS enteros, como 410. No escribas 6:50. El tablero lo muestra '
          'convertido a minutos y segundos, pero la comparación tiene que ser entre números.\n'
          '· Cantidades: números enteros.\n'
          '· Si un agente no tiene medición, deja la celda vacía: sale en rojo y se nota.\n'
          '· Si un indicador queda en blanco para TODO el equipo, el tablero entiende que aún '
          'no se ha medido y no lo cuenta.\n'
          '· La columna CUMPLIMIENTO se calcula sola; es para verlo aquí.', 'p', 2)
    linea('PENDIENTES Y LADO A LADO', 'h2')
    linea('El nombre del agente tiene que quedar escrito IGUAL que en la hoja de indicadores: '
          'por ahí se cruzan las hojas. Usa el desplegable de la primera columna y no habrá '
          'problema.\n\nEn la columna del enlace pega la URL completa de la interacción; el '
          'agente tendrá el botón para abrirla.\n\nEn lado a lado va una fila por '
          'acompañamiento. No se califica con nota: lo que vale es lo que escribas en cada '
          'criterio. Las columnas de criterio son libres.', 'p', 2)
    ws.freeze_panes='A2'

    # ── Indicadores ────────────────────────────────────────────────────
    ws=wb.create_sheet('Indicadores'); ws.sheet_view.showGridLines=False
    fijas=['cedula','nombre_completo','skill','jefe_inmediato']
    ws.cell(1,1,'AGENTES').fill=CAB; ws.cell(1,1).font=BLANCO
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(fijas))
    for k in range(2,len(fijas)+1): ws.cell(1,k).fill=CAB
    for k,n in enumerate(fijas, start=1): ws.cell(2,k,n)

    col=len(fijas)+1; sitio={}
    for i,(nm,tipo,meta,dirn) in enumerate(IND):
        c1=ws.cell(1,col,nm); c1.fill=CAB if i%2==0 else CAB2; c1.font=BLANCO
        c1.alignment=Alignment(horizontal='center', vertical='center')
        ws.merge_cells(start_row=1,start_column=col,end_row=1,end_column=col+2)
        for k in (1,2): ws.cell(1,col+k).fill=CAB if i%2==0 else CAB2
        ws.cell(2,col,'RESULTADO'); ws.cell(2,col+1,'META'); ws.cell(2,col+2,'CUMPLIMIENTO')
        sitio[nm]=(col,tipo,meta,dirn); col+=3
    cp=col
    c1=ws.cell(1,cp,'PRODUCTIVIDAD'); c1.fill=PROD_F; c1.font=BLANCO
    c1.alignment=Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=1,start_column=cp,end_row=1,end_column=cp+len(PROD)-1)
    for k,p in enumerate(PROD):
        ws.cell(1,cp+k).fill=PROD_F; ws.cell(2,cp+k,p)
    ultima=cp+len(PROD)-1
    for c in range(1,ultima+1):
        s=ws.cell(2,c); s.fill=SUAVE; s.font=Font(color=TINTA,bold=True,size=8.5)
        s.alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)
        s.border=BORDE
    ws.row_dimensions[1].height=26; ws.row_dimensions[2].height=30

    for i,(ced,nom,sk,jefe) in enumerate(gente):
        r=3+i
        a=ws.cell(r,1,ced); a.number_format='@'; a.font=Font(size=9.5,color='64748B')
        ws.cell(r,2,nom).font=Font(size=10,bold=True,color=TINTA)
        ws.cell(r,3,sk).font=Font(size=9.5,color='3D566E')
        ws.cell(r,4,jefe).font=Font(size=9,color='7A8CA3')
        for c in range(1,ultima+1):
            cl=ws.cell(r,c); cl.border=BORDE
            if c>len(fijas) and i%2==1: cl.fill=BANDA
        for nm,(c0,tipo,meta,dirn) in sitio.items():
            L=get_column_letter
            res,met,cum=L(c0)+str(r), L(c0+1)+str(r), L(c0+2)+str(r)
            cr,cm,cc=ws[res],ws[met],ws[cum]
            cr.fill=EDITA
            cr.alignment=cm.alignment=cc.alignment=Alignment(horizontal='center')
            cr.number_format=cm.number_format = '0.0%' if tipo=='pct' else '0'
            cm.value=meta; cm.font=Font(size=9.5,color='64748B')
            cc.value='=IF({r}="","",IF({r}{op}{m},"CUMPLE","NO CUMPLE"))'.format(
                r=res, m=met, op='>=' if dirn=='mayor' else '<=')
            cc.fill=CALC; cc.font=Font(size=9,bold=True,color='475569')
        for k in range(len(PROD)):
            cl=ws.cell(r,cp+k); cl.fill=EDITA; cl.number_format='#,##0'
            cl.alignment=Alignment(horizontal='center')
    ws.column_dimensions['A'].width=13; ws.column_dimensions['B'].width=40
    ws.column_dimensions['C'].width=22; ws.column_dimensions['D'].width=30
    for c in range(len(fijas)+1, ultima+1):
        ws.column_dimensions[get_column_letter(c)].width=13
    ws.freeze_panes='E3'

    # ── Listas (oculta) ────────────────────────────────────────────────
    ws=wb.create_sheet('Listas')
    ws.cell(1,1,'Nombres del equipo — no borrar esta hoja').font=Font(bold=True,size=9,color='64748B')
    for i,(ced,nom,sk,jefe) in enumerate(gente):
        ws.cell(2+i,1,nom); ws.cell(2+i,2,ced).number_format='@'
        ws.cell(2+i,3,sk); ws.cell(2+i,4,jefe)
    ws.column_dimensions['A'].width=42
    RANGO='Listas!$A$2:$A$%d' % (1+len(gente))
    ws.sheet_state='hidden'

    # ── Pendientes ─────────────────────────────────────────────────────
    ws=wb.create_sheet('Pendientes_Tipificar'); ws.sheet_view.showGridLines=False
    cols=[('nombre_agente',40),('id_interaccion',34),('link_gestion',46),('cola',14),
          ('macroproceso',24),('proceso',24),('subproceso_uno',24),('subproceso_dos',24),
          ('observaciones',34)]
    for i,(nm,w) in enumerate(cols, start=1):
        c=ws.cell(1,i,nm); c.fill=CAB; c.font=BLANCO; c.border=BORDE
        c.alignment=Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[1].height=28; ws.freeze_panes='A2'
    ws.auto_filter.ref='A1:%s1'%get_column_letter(len(cols))
    dv=DataValidation(type='list', formula1=RANGO, allow_blank=True, errorStyle='warning',
        errorTitle='Nombre fuera del equipo',
        error='Ese nombre no está en la hoja de indicadores. Revisa que esté escrito igual.')
    ws.add_data_validation(dv); dv.add('A2:A3000')
    for r in range(2,500):
        for c in range(1,len(cols)+1):
            cl=ws.cell(r,c); cl.border=BORDE
            if r%2==1: cl.fill=BANDA

    # ── Lado a Lado ────────────────────────────────────────────────────
    ws=wb.create_sheet('Lado a Lado'); ws.sheet_view.showGridLines=False
    cab=['Nombre del agente','Fecha','ID o enlace de la interaccion']+CRIT
    for i,nm in enumerate(cab, start=1):
        c=ws.cell(1,i,nm); c.fill=CAB if i<=3 else CAB2; c.font=BLANCO; c.border=BORDE
        c.alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 40 if i==1 else (
            13 if i==2 else (40 if i==3 else 34))
    ws.row_dimensions[1].height=38; ws.freeze_panes='D2'
    ws.auto_filter.ref='A1:%s1'%get_column_letter(len(cab))
    dv2=DataValidation(type='list', formula1=RANGO, allow_blank=True, errorStyle='warning',
        errorTitle='Nombre fuera del equipo',
        error='Ese nombre no está en la hoja de indicadores. Revisa que esté escrito igual.')
    ws.add_data_validation(dv2); dv2.add('A2:A3000')
    for r in range(2,400):
        for c in range(1,len(cab)+1):
            cl=ws.cell(r,c); cl.border=BORDE
            cl.alignment=Alignment(wrap_text=True, vertical='top')
            if r%2==1: cl.fill=BANDA
            if c==2: cl.number_format='DD/MM/YYYY'

    wb.save(SALIDA+'/'+archivo)
    print('  %-42s %3d agentes · %d supervisor(es)' % (archivo, len(gente), len(jefes)))
    return len(gente)

print('plantillas:')
tot=sum(construir(*e) for e in EQUIPOS)
print('  total %d de 109' % tot)
assert tot==109, tot
